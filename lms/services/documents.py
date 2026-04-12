import io
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from django.utils import timezone

def _set_font(run, size=11, bold=False, italic=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = "Times New Roman"
    if color:
        run.font.color.rgb = RGBColor(*color)

def _add_heading(doc, text, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = doc.add_paragraph()
    p.alignment = align
    run = p.add_run(text)
    _set_font(run, size=size, bold=bold)
    return p

def _add_para(doc, text, size=11, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    run = p.add_run(text)
    _set_font(run, size=size, bold=bold)
    return p

def _add_labeled(doc, label, value, size=11):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_label = p.add_run(f"{label}: ")
    _set_font(run_label, size=size, bold=True)
    run_val = p.add_run(str(value or "-"))
    _set_font(run_val, size=size)
    return p

def _table_set_borders(table):
    from docx.oxml import OxmlElement
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    tblBorders = OxmlElement("w:tblBorders")
    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")
        tblBorders.append(border)
    tblPr.append(tblBorders)
    if tbl.tblPr is None:
        tbl.insert(0, tblPr)

def generate_application_docx(application) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    snap = application.student_snapshot
    faculty_name = snap.faculty_name if snap else ""
    group_name = snap.group_name if snap else ""
    full_name = snap.full_name if snap else "-"

    p_addr = doc.add_paragraph()
    p_addr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    lines = [
        "Buxoro davlat texnika universiteti",
        "rektori S. Siddiqovaga",
        f"{faculty_name or '_____________________'} fakulteti",
        f"{group_name or '_____________________'} guruhi talabasi",
        full_name,
    ]
    for i, line in enumerate(lines):
        run = p_addr.add_run(line)
        _set_font(run, size=11)
        if i < len(lines) - 1:
            p_addr.add_run("\n")

    doc.add_paragraph()
    _add_heading(doc, "ARIZA", size=18, bold=True)
    doc.add_paragraph()

    _add_para(
        doc,
        "Men quyida keltirilgan fanlardan akademik qarzdor bo'lib qoldim. "
        "Ushbu fanlarni belgilangan kredit miqdorini to'lab qayta o'zlashtirishga "
        "ruxsat berishingizni so'rayman.",
        size=11,
    )
    doc.add_paragraph()

    items = application.items.all()
    table = doc.add_table(rows=1, cols=4)
    _table_set_borders(table)
    table.style = "Table Grid"

    hdr_cells = table.rows[0].cells
    headers = ["№", "Fan nomi", "Kredit miqdori", "Semestr"]
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for para in hdr_cells[i].paragraphs:
            for run in para.runs:
                _set_font(run, size=10, bold=True)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for idx, item in enumerate(items, start=1):
        row_cells = table.add_row().cells
        row_cells[0].text = str(idx)
        subject_name = item.subject_snapshot.subject_name if item.subject_snapshot else "-"
        row_cells[1].text = subject_name
        credit = item.subject_snapshot.credit if item.subject_snapshot else 0
        row_cells[2].text = "{:,.2f}".format(credit)
        semester = item.subject_snapshot.semester_name if item.subject_snapshot else "-"
        row_cells[3].text = semester
        for cell in row_cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    _set_font(run, size=10)

    doc.add_paragraph()
    _add_labeled(doc, "Kiritilgan summa", "{:,.2f}".format(application.declared_amount or 0))
    _add_labeled(doc, "Buxgalteriya summasi", "{:,.2f}".format(application.accountant_amount or 0))

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
